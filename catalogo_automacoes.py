#!/usr/bin/env python3
"""
Catálogo de Automações
Organiza as automações da empresa em pastas (setores) com tipo, local e descrição.
A base é um arquivo JSON que pode ficar numa pasta de rede compartilhada.

Requisitos: pip install PySide6
"""
import getpass
import html
import json
import csv
import os
import re
import shutil
import sys
import time
import uuid
from datetime import datetime
from string import Template

from PySide6.QtCore import QRegularExpression, Qt, QTimer, QUrl, Signal
from PySide6.QtGui import (QAction, QActionGroup, QColor, QDesktopServices, QGuiApplication, QIcon, QKeySequence,
                           QPalette, QRegularExpressionValidator, QShortcut, QTextOption)
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFileDialog, QFormLayout,
    QFrame, QHBoxLayout, QHeaderView, QInputDialog, QLabel, QLineEdit, QListWidget,
    QListWidgetItem, QMainWindow, QMenu, QMessageBox, QPlainTextEdit, QPushButton, QSplitter, QStyle, QTabBar, QTextBrowser,
    QTreeWidget, QTreeWidgetItem, QVBoxLayout, QWidget,
)

APP_NAME = "Catálogo de Automações"
BASE_DIR = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")          # base padrão distribuída junto com o app
PREFS_DIR = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "CatalogoAutomacoes")
PREFS_PATH = os.path.join(PREFS_DIR, "prefs.json")            # preferências de cada usuário
DEFAULT_TYPES = ["Python", "Power Automate", "Macro VBA", "RPA", "SQL / Procedure",
                 "Script BAT / PowerShell", "Power BI", "Agendador de Tarefas", "Outro"]
FIELD_LABELS = {"name": "Nome", "type": "Tipo", "companies": "Empresas", "folder_id": "Pasta",
                "location": "Local", "description": "Descrição"}
HISTORY_LIMIT = 5000
NO_COMPANY = "\x00sem-empresa"   # marcador interno para "automações sem empresa"
CODE_RE = re.compile(r"[0-9]{4}")
LABEL_RE = re.compile(r"^\s*([0-9]{1,4})(?:\s+(.*?))?\s*$")


def norm_code(value):
    """Converte 10, 10.0, '10', ' 0010 ' em '0010'. Levanta ValueError se não for um código de até 4 dígitos."""
    if isinstance(value, bool) or value is None:
        raise ValueError
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError
        value = int(value)
    text = str(value).strip()
    if not re.fullmatch(r"[0-9]+", text) or int(text) > 9999:
        raise ValueError
    return str(int(text)).zfill(4)


def split_label(text):
    """'10 EMPRESA X' -> ('0010', 'EMPRESA X'); '10' -> ('0010', ''); sem código -> (None, texto)."""
    text = " ".join(str(text).split())
    m = LABEL_RE.match(text)
    if m:
        return m.group(1).zfill(4), (m.group(2) or "")
    return None, text
MAX_RECENT = 8
ROLE_ID = Qt.UserRole
ROLE_KIND = Qt.UserRole + 1
ICON_FILE = "icone.ico"


def resource(name):
    """Caminho de um arquivo embutido no .exe (PyInstaller) ou ao lado do .py."""
    return os.path.join(getattr(sys, "_MEIPASS", BASE_DIR), name)


def now():
    return datetime.now().strftime("%d/%m/%Y %H:%M")


def now_full():
    return datetime.now().strftime("%d/%m/%Y %H:%M:%S")


def human_size(n):
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}".replace(".", ",")
        n /= 1024


def current_user():
    try:
        return getpass.getuser()
    except Exception:
        return "desconhecido"


# --------------------------------------------------------------------------- base
class Database:
    """Base em JSON com trava de arquivo, segura para vários usuários na rede."""
    LOCK_TIMEOUT = 10
    STALE_LOCK = 30

    def __init__(self, path):
        self.path = os.path.abspath(path)
        self.lock_path = self.path + ".lock"
        self.attach_root = os.path.join(os.path.dirname(self.path), f"{self.name}_anexos")
        self.data = {"folders": [], "items": [], "history": [], "companies": []}
        self.mtime = None
        if os.path.exists(self.path):
            self.load()
            if self.migration_note:  # grava já convertida, para todos verem no formato novo
                note = self.migration_note
                self.mutate(lambda d: self._log(d, "converteu", "base", None, os.path.basename(self.path), note=note))
        else:
            self.mutate(lambda d: self._log(d, "criou", "base", None, os.path.basename(self.path)))

    @property
    def name(self):
        return os.path.splitext(os.path.basename(self.path))[0]

    def load(self):
        with open(self.path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict) or "folders" not in data or "items" not in data:
            raise ValueError("O arquivo escolhido não é uma base do Catálogo de Automações.")
        data.setdefault("history", [])
        data.setdefault("companies", [])
        self.migration_note = self._migrate(data)
        self.data = data
        self.mtime = os.path.getmtime(self.path)

    def changed_on_disk(self):
        try:
            return os.path.getmtime(self.path) != self.mtime
        except OSError:
            return False

    def _lock(self):
        start = time.time()
        while True:
            try:
                fd = os.open(self.lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.write(fd, f"{current_user()} {now()}".encode())
                os.close(fd)
                return
            except FileExistsError:
                try:
                    if time.time() - os.path.getmtime(self.lock_path) > self.STALE_LOCK:
                        os.remove(self.lock_path)
                        continue
                except OSError:
                    pass
                if time.time() - start > self.LOCK_TIMEOUT:
                    raise RuntimeError("A base está ocupada por outro usuário. Tente novamente em instantes.")
                time.sleep(0.15)

    def _unlock(self):
        try:
            os.remove(self.lock_path)
        except OSError:
            pass

    def mutate(self, fn):
        """Recarrega a versão mais recente, aplica a alteração e salva de forma atômica."""
        self._lock()
        try:
            if os.path.exists(self.path):
                self.load()
            result = fn(self.data)
            tmp = f"{self.path}.{uuid.uuid4().hex}.tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(self.data, f, ensure_ascii=False, indent=2)
            if os.path.exists(self.path):
                shutil.copy2(self.path, self.path + ".bak")
            os.replace(tmp, self.path)
            self.mtime = os.path.getmtime(self.path)
            return result
        except Exception:
            if os.path.exists(self.path):
                try:
                    self.load()
                except Exception:
                    pass
            raise
        finally:
            self._unlock()

    # histórico
    @staticmethod
    def _log(d, action, kind, oid, name, path="", changes=None, note=""):
        hist = d.setdefault("history", [])
        hist.append({"id": uuid.uuid4().hex, "at": now_full(), "user": current_user(),
                     "action": action, "kind": kind, "target_id": oid, "name": name,
                     "path": path, "changes": changes or [], "note": note})
        if len(hist) > HISTORY_LIMIT:
            del hist[:len(hist) - HISTORY_LIMIT]

    def _fmt(self, key, value):
        if key == "folder_id":
            return self.folder_path(value or None)
        if key == "companies":
            return self._labels(value)
        return value or ""

    def _snapshot(self, it, empty_side="old"):
        rows = []
        for k, label in FIELD_LABELS.items():
            val = self._fmt(k, it.get(k))
            if val:
                rows.append([label, "", val] if empty_side == "old" else [label, val, ""])
        return rows

    # consultas
    def folder(self, fid):
        return next((f for f in self.data["folders"] if f["id"] == fid), None)

    def item(self, iid):
        return next((i for i in self.data["items"] if i["id"] == iid), None)

    def subfolders(self, parent):
        return sorted((f for f in self.data["folders"] if f.get("parent_id") == parent),
                      key=lambda f: f["name"].lower())

    def items_in(self, fid):
        return sorted((i for i in self.data["items"] if i.get("folder_id") == fid),
                      key=lambda i: i["name"].lower())

    def folder_path(self, fid):
        parts = []
        while fid:
            f = self.folder(fid)
            if not f:
                break
            parts.insert(0, f["name"])
            fid = f.get("parent_id")
        return " / ".join(["Início"] + parts)

    def folder_paths(self):
        out = []

        def walk(parent):
            for f in self.subfolders(parent):
                out.append((f["id"], self.folder_path(f["id"])))
                walk(f["id"])
        walk(None)
        return out

    @staticmethod
    def descendants(data, fid):
        ids, stack = {fid}, [fid]
        while stack:
            cur = stack.pop()
            for f in data["folders"]:
                if f.get("parent_id") == cur and f["id"] not in ids:
                    ids.add(f["id"])
                    stack.append(f["id"])
        return ids

    # empresas: cadastro {"code": "0001", "name": "..."}; as automações guardam só os códigos
    @staticmethod
    def _migrate(data):
        """Converte empresas do formato antigo (texto "0001 NOME") para código + nome.
        Devolve um resumo do que foi feito ou "" se a base já estava no formato novo."""
        comps, items = data.get("companies", []), data["items"]
        legacy = any(not isinstance(c, dict) for c in comps) or any(
            not (isinstance(c, str) and CODE_RE.fullmatch(c)) for i in items for c in i.get("companies") or [])
        if not legacy:
            return ""
        reg, by_text, raw = {}, {}, 0

        def learn(entry):
            if isinstance(entry, dict):
                code, name = entry.get("code"), " ".join(str(entry.get("name", "")).split())
                try:
                    code = norm_code(code)
                except ValueError:
                    return None
            else:
                code, name = split_label(entry)
                if code is None:
                    return None
            if not reg.get(code):  # mantém o primeiro nome não vazio
                reg[code] = name
            if name:
                by_text.setdefault(name.lower(), code)
            return code

        for c in comps:
            raw += 1
            learn(c)
        for i in items:
            for c in i.get("companies") or []:
                learn(c)
        # textos sem código ganham um código livre, do 9999 para baixo
        pending = []
        for c in list(comps) + [c for i in items for c in i.get("companies") or []]:
            if not isinstance(c, dict) and split_label(c)[0] is None:
                key = " ".join(str(c).split()).lower()
                if key and key not in by_text:
                    by_text[key] = None
                    pending.append(" ".join(str(c).split()))
        free = (str(n).zfill(4) for n in range(9999, 0, -1) if str(n).zfill(4) not in reg)
        provisional = []
        for name in pending:
            code = next(free)
            reg[code] = name
            by_text[name.lower()] = code
            provisional.append(f"{code} {name}")

        def code_of(entry):
            if isinstance(entry, dict):
                try:
                    return norm_code(entry.get("code"))
                except ValueError:
                    return None
            code, _ = split_label(entry)
            return code or by_text.get(" ".join(str(entry).split()).lower())

        for i in items:
            codes = {code_of(c) for c in i.get("companies") or []}
            i["companies"] = sorted(c for c in codes if c)
        data["companies"] = [{"code": c, "name": reg[c]} for c in sorted(reg)]
        note = f"Empresas convertidas para código + nome: {raw} registro(s) viraram {len(reg)} empresa(s)."
        if raw > len(reg):
            note += f" {raw - len(reg)} duplicata(s) com o mesmo código foram unificadas."
        if provisional:
            note += " Empresas sem código receberam código provisório: " + "; ".join(provisional) + "."
        return note

    def companies(self):
        return sorted(self.data["companies"], key=lambda c: c["code"])

    def company(self, code):
        return next((c for c in self.data["companies"] if c["code"] == code), None)

    def company_label(self, code):
        c = self.company(code)
        return f'{code} {c["name"]}'.strip() if c else code

    def company_items(self, code):
        """Automações de uma empresa; code=None devolve as sem empresa."""
        if code is None:
            lst = [i for i in self.data["items"] if not i.get("companies")]
        else:
            lst = [i for i in self.data["items"] if code in (i.get("companies") or [])]
        return sorted(lst, key=lambda i: i["name"].lower())

    def _labels(self, codes):
        return ", ".join(self.company_label(c) for c in codes or [])

    def _register(self, d, pairs):
        """Cadastra empresas novas vindas da tela de atribuição; ignora códigos já existentes."""
        for code, name in pairs or []:
            code, name = norm_code(code), " ".join(name.split())
            if name and not self.company(code):
                d["companies"].append({"code": code, "name": name})
                self._log(d, "criou", "empresa", None, f"{code} {name}")

    def _set_companies(self, d, it, new):
        old = it.get("companies") or []
        new = sorted(set(new))
        if old == new:
            return 0
        self._log(d, "editou", "automação", it["id"], it["name"], self.folder_path(it.get("folder_id")),
                  [["Empresas", self._labels(old), self._labels(new)]])
        it["companies"] = new
        it["updated_at"], it["updated_by"] = now(), current_user()
        return 1

    def add_company(self, code, name):
        code, name = norm_code(code), " ".join(name.split())
        if not name:
            raise ValueError("Informe o nome da empresa.")

        def op(d):
            if self.company(code):
                raise ValueError(f'O código {code} já está cadastrado para "{self.company_label(code)}".')
            d["companies"].append({"code": code, "name": name})
            self._log(d, "criou", "empresa", None, f"{code} {name}")
            return code
        return self.mutate(op)

    def edit_company(self, code, new_code, new_name):
        """Altera código e/ou nome. Trocar o código atualiza todas as automações ligadas."""
        new_code, new_name = norm_code(new_code), " ".join(new_name.split())
        if not new_name:
            raise ValueError("Informe o nome da empresa.")

        def op(d):
            c = self.company(code)
            if not c:
                raise RuntimeError("Esta empresa foi excluída por outro usuário.")
            if new_code != code and self.company(new_code):
                raise ValueError(f'O código {new_code} já está cadastrado para "{self.company_label(new_code)}".')
            changes = []
            if new_code != code:
                changes.append(["Código", code, new_code])
            if new_name != c["name"]:
                changes.append(["Nome", c["name"], new_name])
            if not changes:
                return code
            c["code"], c["name"] = new_code, new_name
            if new_code != code:
                for it in d["items"]:
                    if code in (it.get("companies") or []):
                        it["companies"] = sorted(new_code if x == code else x for x in it["companies"])
            self._log(d, "editou", "empresa", None, f"{new_code} {new_name}", changes=changes)
            return new_code
        return self.mutate(op)

    def remove_company(self, code):
        def op(d):
            label = self.company_label(code)
            n = 0
            for it in d["items"]:
                if code in (it.get("companies") or []):
                    n += self._set_companies(d, it, [x for x in it["companies"] if x != code])
            d["companies"] = [c for c in d["companies"] if c["code"] != code]
            self._log(d, "excluiu", "empresa", None, label,
                      note=f"Removida de {n} automação(ões). As automações não foram excluídas.")
            return n
        return self.mutate(op)

    def set_item_companies(self, iid, codes, new_companies=None):
        def op(d):
            it = self.item(iid)
            if not it:
                raise RuntimeError("Esta automação foi excluída por outro usuário.")
            self._register(d, new_companies)
            return self._set_companies(d, it, [c for c in codes if self.company(c)])
        return self.mutate(op)

    def set_company_items(self, code, ids):
        """Deixa a empresa ligada exatamente às automações de ids. Devolve quantas mudaram."""
        ids = set(ids)

        def op(d):
            if not self.company(code):
                raise RuntimeError("Esta empresa foi excluída por outro usuário.")
            n = 0
            for it in d["items"]:
                comps = it.get("companies") or []
                if code in comps and it["id"] not in ids:
                    n += self._set_companies(d, it, [c for c in comps if c != code])
                elif code not in comps and it["id"] in ids:
                    n += self._set_companies(d, it, comps + [code])
            return n
        return self.mutate(op)

    def all_types(self):
        used = {i.get("type", "") for i in self.data["items"] if i.get("type")}
        return DEFAULT_TYPES + sorted(used - set(DEFAULT_TYPES), key=str.lower)

    # alterações
    @staticmethod
    def _check_name(data, name, parent, ignore_id=None):
        for f in data["folders"]:
            if f.get("parent_id") == parent and f["id"] != ignore_id and f["name"].lower() == name.lower():
                raise ValueError(f'Já existe uma pasta chamada "{name}" aqui.')

    def add_folder(self, name, parent):
        def op(d):
            if parent and not self.folder(parent):
                raise RuntimeError("A pasta de destino foi excluída por outro usuário.")
            self._check_name(d, name, parent)
            fid = uuid.uuid4().hex
            d["folders"].append({"id": fid, "name": name, "parent_id": parent,
                                 "created_at": now(), "created_by": current_user()})
            self._log(d, "criou", "pasta", fid, name, self.folder_path(parent))
            return fid
        return self.mutate(op)

    def rename_folder(self, fid, name):
        def op(d):
            f = self.folder(fid)
            if not f:
                raise RuntimeError("Esta pasta foi excluída por outro usuário.")
            self._check_name(d, name, f.get("parent_id"), ignore_id=fid)
            old = f["name"]
            f["name"] = name
            self._log(d, "renomeou", "pasta", fid, name, self.folder_path(f.get("parent_id")),
                      [["Nome", old, name]])
        self.mutate(op)

    def delete_folder(self, fid):
        def op(d):
            f = self.folder(fid)
            if not f:
                return
            ids = self.descendants(d, fid)
            gone = [i for i in d["items"] if i.get("folder_id") in ids]
            for it in gone:
                self._log(d, "excluiu", "automação", it["id"], it["name"],
                          self.folder_path(it.get("folder_id")), self._snapshot(it, "new"),
                          f'Excluída junto com a pasta "{f["name"]}".')
            self._log(d, "excluiu", "pasta", fid, f["name"], self.folder_path(f.get("parent_id")),
                      note=f"{len(ids) - 1} subpasta(s) e {len(gone)} automação(ões) excluídas junto.")
            d["folders"] = [x for x in d["folders"] if x["id"] not in ids]
            d["items"] = [i for i in d["items"] if i.get("folder_id") not in ids]
            return [i["id"] for i in gone]
        for iid in self.mutate(op) or []:
            self._trash(iid)

    def save_item(self, values, iid=None, new_companies=None):
        def op(d):
            self._register(d, new_companies)
            values["companies"] = sorted({c for c in values.get("companies") or [] if self.company(c)})
            fid = values.get("folder_id")
            if fid and not self.folder(fid):
                raise RuntimeError("A pasta escolhida foi excluída por outro usuário.")
            stamp = {"updated_at": now(), "updated_by": current_user()}
            if iid:
                it = self.item(iid)
                if not it:
                    raise RuntimeError("Esta automação foi excluída por outro usuário.")
                changes = []
                for k, label in FIELD_LABELS.items():
                    old, new = self._fmt(k, it.get(k)), self._fmt(k, values.get(k))
                    if old != new:
                        changes.append([label, old, new])
                if changes:
                    it.update(values, **stamp)
                    self._log(d, "editou", "automação", iid, it["name"],
                              self.folder_path(it.get("folder_id")), changes)
                return iid
            new_id = uuid.uuid4().hex
            it = {"id": new_id, **values, "created_at": now(), "created_by": current_user(), **stamp}
            d["items"].append(it)
            self._log(d, "criou", "automação", new_id, it["name"],
                      self.folder_path(fid), self._snapshot(it, "old"))
            return new_id
        return self.mutate(op)

    def delete_item(self, iid):
        def op(d):
            it = self.item(iid)
            if not it:
                return
            self._log(d, "excluiu", "automação", iid, it["name"],
                      self.folder_path(it.get("folder_id")), self._snapshot(it, "new"))
            d["items"] = [i for i in d["items"] if i["id"] != iid]
            return True
        if self.mutate(op):
            self._trash(iid)

    # anexos
    def attachment_path(self, iid, att):
        return os.path.join(self.attach_root, iid, att["file"])

    def _trash(self, iid, file=None):
        """Move anexos de itens excluídos para _excluidos em vez de apagar."""
        src = os.path.join(self.attach_root, iid, file) if file else os.path.join(self.attach_root, iid)
        if not os.path.exists(src):
            return
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dst = os.path.join(self.attach_root, "_excluidos", iid, f"{stamp}_{file}") if file \
            else os.path.join(self.attach_root, "_excluidos", f"{iid}_{stamp}")
        try:
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            shutil.move(src, dst)
        except OSError:
            pass

    def add_attachments(self, iid, paths):
        if not self.item(iid):
            raise RuntimeError("Esta automação foi excluída por outro usuário.")
        folder = os.path.join(self.attach_root, iid)
        os.makedirs(folder, exist_ok=True)
        copied = []
        try:
            for src in paths:
                base, ext = os.path.splitext(os.path.basename(src))
                name, n = base + ext, 2
                while os.path.exists(os.path.join(folder, name)):
                    name, n = f"{base} ({n}){ext}", n + 1
                shutil.copy2(src, os.path.join(folder, name))
                copied.append({"id": uuid.uuid4().hex, "file": name,
                               "size": os.path.getsize(os.path.join(folder, name)),
                               "added_at": now(), "added_by": current_user()})
        except OSError as ex:
            for c in copied:
                self._remove_file(os.path.join(folder, c["file"]))
            raise RuntimeError(f"Não foi possível copiar o arquivo para a pasta de anexos:\n{ex}")

        def op(d):
            it = self.item(iid)
            if not it:
                raise RuntimeError("Esta automação foi excluída por outro usuário.")
            it.setdefault("attachments", []).extend(copied)
            for c in copied:
                self._log(d, "anexou", "automação", iid, it["name"], self.folder_path(it.get("folder_id")),
                          [["Arquivo", "", c["file"]]])
        try:
            self.mutate(op)
        except Exception:
            for c in copied:
                self._remove_file(os.path.join(folder, c["file"]))
            try:
                os.rmdir(folder)  # só remove se ficou vazia
            except OSError:
                pass
            raise
        return len(copied)

    @staticmethod
    def _remove_file(path):
        try:
            os.remove(path)
        except OSError:
            pass

    def remove_attachment(self, iid, att_id):
        def op(d):
            it = self.item(iid)
            att = next((a for a in (it or {}).get("attachments", []) if a["id"] == att_id), None)
            if not att:
                return None
            it["attachments"] = [a for a in it["attachments"] if a["id"] != att_id]
            self._log(d, "desanexou", "automação", iid, it["name"], self.folder_path(it.get("folder_id")),
                      [["Arquivo", att["file"], ""]])
            return att["file"]
        file = self.mutate(op)
        if file:
            self._trash(iid, file)


# --------------------------------------------------------------------------- temas
THEMES = {
    "Claro": dict(bg="#f3f4f6", surface="#ffffff", border="#e5e7eb", input="#d1d5db", text="#1f2937",
                  muted="#6b7280", faint="#9ca3af", accent="#2563eb", accent_hover="#1d4ed8",
                  accent_text="#ffffff", sel="#dbeafe", sel_text="#1e3a8a", hover="#f1f5ff",
                  danger="#dc2626", danger_bg="#fef2f2", btn_hover="#f9fafb", add="#15803d"),
    "Escuro": dict(bg="#0f1115", surface="#181b21", border="#2a2f3a", input="#3a404c", text="#e5e7eb",
                   muted="#9ca3af", faint="#6b7280", accent="#3b82f6", accent_hover="#2563eb",
                   accent_text="#ffffff", sel="#1e3a5f", sel_text="#e0ecff", hover="#222834",
                   danger="#f87171", danger_bg="#3a1d1d", btn_hover="#222834", add="#4ade80"),
    "Grafite": dict(bg="#1c1f24", surface="#24282e", border="#353a42", input="#4a505a", text="#e6e6e6",
                    muted="#a0a7b1", faint="#767d88", accent="#14b8a6", accent_hover="#0d9488",
                    accent_text="#06201d", sel="#134e4a", sel_text="#ccfbf1", hover="#2d323a",
                    danger="#f87171", danger_bg="#3b2222", btn_hover="#2d323a", add="#5eead4"),
    "Verde": dict(bg="#f2f5f3", surface="#ffffff", border="#e2e8e4", input="#cfd8d2", text="#1c2a22",
                  muted="#5f6f66", faint="#94a39a", accent="#059669", accent_hover="#047857",
                  accent_text="#ffffff", sel="#d1fae5", sel_text="#064e3b", hover="#ecfdf5",
                  danger="#dc2626", danger_bg="#fef2f2", btn_hover="#f6f8f7", add="#047857"),
}

STYLE = Template("""
* { font-family: "Segoe UI", "Inter", sans-serif; font-size: 10pt; color: $text; }
QMainWindow, #central, QDialog, QMessageBox, QInputDialog { background: $bg; }
#topbar { background: $surface; border-bottom: 1px solid $border; }
#crumb { font-size: 12pt; font-weight: 600; padding-left: 6px; }
#tabrow { background: $surface; }
QTabBar#maintabs { background: transparent; }
QTabBar#maintabs::tab { background: transparent; border: none; border-bottom: 2px solid transparent;
    padding: 10px 20px 8px 20px; margin-right: 4px; color: $muted; font-size: 10.5pt; font-weight: 600; }
QTabBar#maintabs::tab:selected { color: $text; border-bottom: 2px solid $accent; }
QTabBar#maintabs::tab:hover:!selected { color: $text; background: $hover; }
#section { color: $muted; font-size: 8pt; font-weight: 600; letter-spacing: 1px; }
QListWidget { background: $bg; border: 1px dashed $input; border-radius: 6px; padding: 4px; outline: none; }
QListWidget::item { padding: 4px; border-radius: 4px; }
QListWidget::item:hover { background: $hover; }
QListWidget::item:selected { background: $sel; color: $sel_text; }
QListWidget::indicator { width: 15px; height: 15px; }
QMenuBar { background: $surface; border-bottom: 1px solid $border; padding: 2px 6px; }
QMenuBar::item { background: transparent; padding: 5px 10px; border-radius: 4px; }
QMenuBar::item:selected { background: $hover; }
QMenu { background: $surface; border: 1px solid $border; padding: 4px; }
QMenu::item { padding: 6px 22px; border-radius: 4px; }
QMenu::item:selected { background: $sel; color: $sel_text; }
QMenu::item:disabled { color: $faint; }
QMenu::separator { height: 1px; background: $border; margin: 4px 8px; }
QTreeWidget { background: $surface; border: 1px solid $border; border-radius: 8px; padding: 4px; outline: none; }
QTreeWidget::item { padding: 6px 4px; border-radius: 5px; }
QTreeWidget::item:hover { background: $hover; }
QTreeWidget::item:selected { background: $sel; color: $sel_text; }
QHeaderView::section { background: $surface; border: none; border-bottom: 1px solid $border;
    padding: 8px 6px; font-weight: 600; color: $muted; }
QLineEdit, QComboBox, QPlainTextEdit { background: $surface; border: 1px solid $input;
    border-radius: 6px; padding: 6px 8px; selection-background-color: $accent; selection-color: $accent_text; }
QLineEdit:focus, QComboBox:focus, QPlainTextEdit:focus { border: 1px solid $accent; }
QComboBox QAbstractItemView { background: $surface; border: 1px solid $border;
    selection-background-color: $sel; selection-color: $sel_text; }
QPushButton { background: $surface; border: 1px solid $input; border-radius: 6px; padding: 7px 14px; }
QPushButton:hover { background: $btn_hover; border-color: $faint; }
QPushButton:disabled { color: $faint; }
QPushButton[primary="true"] { background: $accent; color: $accent_text; border: 1px solid $accent; font-weight: 600; }
QPushButton[primary="true"]:hover { background: $accent_hover; }
QPushButton[danger="true"] { color: $danger; }
QPushButton[danger="true"]:hover { background: $danger_bg; border-color: $danger; }
QPushButton[danger="true"]:disabled { color: $faint; background: $surface; border-color: $input; }
QPushButton#link { border: none; background: transparent; color: $accent; padding: 2px 8px; }
#details, #panel { background: $surface; border: 1px solid $border; border-radius: 8px; }
QTextBrowser { border: none; background: transparent; }
QSplitter::handle { background: transparent; width: 8px; height: 8px; }
QStatusBar { background: $surface; border-top: 1px solid $border; }
QStatusBar QLabel { color: $muted; }
QScrollBar:vertical { background: transparent; width: 10px; margin: 2px; }
QScrollBar:horizontal { background: transparent; height: 10px; margin: 2px; }
QScrollBar::handle { background: $input; border-radius: 4px; min-height: 24px; min-width: 24px; }
QScrollBar::handle:hover { background: $faint; }
QScrollBar::add-line, QScrollBar::sub-line { width: 0; height: 0; }
QScrollBar::add-page, QScrollBar::sub-page { background: transparent; }
QToolTip { background: $surface; color: $text; border: 1px solid $border; padding: 4px; }
""")


def apply_theme(app, name):
    t = THEMES.get(name) or THEMES["Claro"]
    pal = QPalette()
    roles = {QPalette.Window: "bg", QPalette.WindowText: "text", QPalette.Base: "surface",
             QPalette.AlternateBase: "bg", QPalette.Text: "text", QPalette.Button: "surface",
             QPalette.ButtonText: "text", QPalette.Highlight: "accent",
             QPalette.HighlightedText: "accent_text", QPalette.ToolTipBase: "surface",
             QPalette.ToolTipText: "text", QPalette.PlaceholderText: "faint", QPalette.Link: "accent"}
    for role, key in roles.items():
        pal.setColor(role, QColor(t[key]))
    pal.setColor(QPalette.Disabled, QPalette.Text, QColor(t["faint"]))
    pal.setColor(QPalette.Disabled, QPalette.ButtonText, QColor(t["faint"]))
    app.setPalette(pal)
    app.setStyleSheet(STYLE.substitute(t))
    return t


# --------------------------------------------------------------------------- preferências
def load_json(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_prefs(prefs):
    try:
        os.makedirs(PREFS_DIR, exist_ok=True)
        with open(PREFS_PATH, "w", encoding="utf-8") as f:
            json.dump(prefs, f, ensure_ascii=False, indent=2)
    except OSError:
        pass


def remember_base(prefs, path):
    path = os.path.abspath(path)
    key = os.path.normcase(path)
    recent = [p for p in prefs.get("recent", []) if os.path.normcase(p) != key]
    prefs["recent"] = [path] + recent[:MAX_RECENT - 1]
    prefs["db_path"] = path
    save_prefs(prefs)


FILTER = "Base de automações (*.json)"


def ask_open_base(parent):
    return QFileDialog.getOpenFileName(parent, "Abrir base existente", "", FILTER)[0] or None


def ask_new_base(parent):
    path = QFileDialog.getSaveFileName(parent, "Criar nova base", "automacoes.json", FILTER,
                                       options=QFileDialog.DontConfirmOverwrite)[0]
    if not path:
        return None
    if not path.lower().endswith(".json"):
        path += ".json"
    if os.path.exists(path):
        ans = QMessageBox.question(parent, APP_NAME, "Já existe uma base nesse local.\n\nDeseja abrir essa base?")
        return path if ans == QMessageBox.Yes else None
    return path


def ask_base_startup():
    box = QMessageBox()
    box.setWindowTitle(APP_NAME)
    box.setText("Onde fica a base de automações?")
    box.setInformativeText("Escolha um arquivo .json numa pasta de rede que todos acessem.")
    b_open = box.addButton("Abrir base existente", QMessageBox.AcceptRole)
    b_new = box.addButton("Criar nova base", QMessageBox.ActionRole)
    box.addButton("Cancelar", QMessageBox.RejectRole)
    box.exec()
    if box.clickedButton() == b_open:
        return ask_open_base(None)
    if box.clickedButton() == b_new:
        return ask_new_base(None)
    return None


# --------------------------------------------------------------------------- diálogos
class CompanyDialog(QDialog):
    """Cadastro ou edição de empresa: código de até 4 dígitos + nome."""

    def __init__(self, parent, code="", name=""):
        super().__init__(parent)
        self.setWindowTitle("Editar empresa" if code else "Nova empresa")
        self.setMinimumWidth(480)
        self.code = QLineEdit(code)
        self.code.setValidator(QRegularExpressionValidator(QRegularExpression(r"[0-9]{0,4}")))
        self.code.setPlaceholderText("Ex.: 10 vira 0010")
        self.code.setMaximumWidth(140)
        self.code.editingFinished.connect(self._pad)
        self.name = QLineEdit(name)
        self.name.setPlaceholderText("Razão social ou nome usado no escritório")
        form = QFormLayout()
        form.setSpacing(10)
        form.addRow("Código *", self.code)
        form.addRow("Nome *", self.name)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("Salvar")
        buttons.button(QDialogButtonBox.Save).setProperty("primary", True)
        buttons.button(QDialogButtonBox.Cancel).setText("Cancelar")
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(20, 20, 20, 16)
        lay.addLayout(form)
        lay.addSpacing(8)
        lay.addWidget(buttons)

    def _pad(self):
        if self.code.text().strip():
            self.code.setText(norm_code(self.code.text()))

    def _accept(self):
        if not self.code.text().strip():
            QMessageBox.warning(self, APP_NAME, "Informe o código da empresa.")
            self.code.setFocus()
            return
        if not self.name.text().strip():
            QMessageBox.warning(self, APP_NAME, "Informe o nome da empresa.")
            self.name.setFocus()
            return
        self._pad()
        self.accept()

    def values(self):
        return norm_code(self.code.text()), " ".join(self.name.text().split())


class ImportCodesDialog(QDialog):
    """Escolhe planilha, aba e coluna de onde ler os códigos das empresas."""

    def __init__(self, parent, path):
        super().__init__(parent)
        self.path = path
        self.setWindowTitle("Importar empresas de planilha")
        self.setMinimumWidth(520)
        self.wb = None
        self.rows = []
        self.sheet = QComboBox()
        self.column = QComboBox()
        self.header = QCheckBox("A primeira linha é cabeçalho")
        self.replace = QCheckBox("Desmarcar as empresas já marcadas antes de importar")
        self.preview = QLabel()
        self.preview.setWordWrap(True)
        self.preview.setObjectName("section")
        form = QFormLayout()
        form.setSpacing(10)
        form.addRow("Arquivo", QLabel(os.path.basename(path)))
        form.addRow("Aba", self.sheet)
        form.addRow("Coluna dos códigos", self.column)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("Importar")
        buttons.button(QDialogButtonBox.Ok).setProperty("primary", True)
        buttons.button(QDialogButtonBox.Cancel).setText("Cancelar")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(20, 20, 20, 16)
        lay.addLayout(form)
        lay.addWidget(self.header)
        lay.addWidget(self.replace)
        lay.addWidget(self.preview)
        lay.addSpacing(8)
        lay.addWidget(buttons)

        self._open()
        self.sheet.currentIndexChanged.connect(self._load_sheet)
        self.column.currentIndexChanged.connect(self._guess_header)
        self.header.toggled.connect(self._update_preview)
        self._load_sheet()

    def _open(self):
        ext = os.path.splitext(self.path)[1].lower()
        if ext == ".csv":
            self.sheet.addItem("CSV")
            return
        if ext == ".xls":
            raise RuntimeError("Arquivos .xls (Excel 97-2003) não são suportados. Abra no Excel e salve como .xlsx.")
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError("O componente de leitura de Excel (openpyxl) não está instalado.\n"
                               "Rode: pip install openpyxl")
        self.wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        self.sheet.addItems(self.wb.sheetnames)

    def _read_csv(self):
        for enc in ("utf-8-sig", "cp1252"):
            try:
                with open(self.path, newline="", encoding=enc) as f:
                    text = f.read()
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        return [list(r) for r in csv.reader(text.splitlines(), dialect)]

    def _load_sheet(self, *_):
        if self.wb is None:
            self.rows = self._read_csv()
        else:
            ws = self.wb[self.sheet.currentText()]
            self.rows = [list(r) for r in ws.iter_rows(values_only=True)]
        while self.rows and all(v in (None, "") for v in self.rows[-1]):
            self.rows.pop()
        ncols = max((len(r) for r in self.rows), default=0)
        first = self.rows[0] if self.rows else []
        self.column.blockSignals(True)
        self.column.clear()
        best, best_score = 0, -1
        for c in range(ncols):
            letter = self._letter(c)
            head = first[c] if c < len(first) else None
            self.column.addItem(f"{letter}  ({head})" if head not in (None, "") else letter, c)
            score = sum(1 for r in self.rows[:200] if c < len(r) and self._code(r[c]))
            if head and "cod" in str(head).lower():
                score += 1000
            if score > best_score:
                best, best_score = c, score
        self.column.setCurrentIndex(best)
        self.column.blockSignals(False)
        self._guess_header()

    @staticmethod
    def _letter(i):
        s = ""
        i += 1
        while i:
            i, r = divmod(i - 1, 26)
            s = chr(65 + r) + s
        return s

    @staticmethod
    def _code(value):
        """Código de um valor de célula: número, '10', '0010' ou '0010 NOME DA EMPRESA'."""
        if value is None or value == "":
            return None
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            try:
                return norm_code(value)
            except ValueError:
                return None
        code, _ = split_label(value)
        return code

    def _guess_header(self, *_):
        c = self.column.currentData()
        first = self.rows[0] if self.rows else []
        is_head = c is not None and c < len(first) and first[c] not in (None, "") and not self._code(first[c])
        self.header.blockSignals(True)
        self.header.setChecked(is_head)
        self.header.blockSignals(False)
        self._update_preview()

    def _values(self):
        c = self.column.currentData()
        if c is None:
            return []
        rows = self.rows[1:] if self.header.isChecked() else self.rows
        return [r[c] for r in rows if c < len(r) and r[c] not in (None, "")]

    def _update_preview(self, *_):
        vals = self._values()
        ok = sum(1 for v in vals if self._code(v))
        self.preview.setText(f"{len(vals)} célula(s) preenchida(s) na coluna, {ok} com código válido.")

    def result_codes(self):
        """(códigos válidos sem repetição, valores que não são código)."""
        codes, invalid = [], []
        for v in self._values():
            code = self._code(v)
            if code:
                if code not in codes:
                    codes.append(code)
            else:
                invalid.append(str(v))
        return codes, invalid

    def done(self, r):
        if self.wb is not None:
            self.wb.close()
        super().done(r)


class CompanyPicker(QWidget):
    """Lista marcável de empresas (código + nome) com filtro, cadastro rápido e importação de planilha."""

    def __init__(self, db, selected, list_height=130):
        super().__init__()
        self.db = db
        self.pending = {}  # empresas novas digitadas aqui: código -> nome
        self.edit = QLineEdit()
        self.edit.setPlaceholderText("Filtrar, ou digitar \"código nome\" e Enter para cadastrar")
        self.edit.textChanged.connect(self._filter)
        self.edit.returnPressed.connect(self.add)
        btn = QPushButton("Adicionar")
        btn.clicked.connect(self.add)
        self.btn_import = QPushButton("Importar do Excel...")
        self.btn_import.clicked.connect(self.import_file)
        self.list = QListWidget()
        self.list.setFixedHeight(list_height)
        chosen = set(selected or [])
        self.list.setUniformItemSizes(True)
        for c in db.companies():
            self._insert(c["code"], c["name"], c["code"] in chosen)
        row = QHBoxLayout()
        row.setContentsMargins(0, 0, 0, 0)
        row.addWidget(self.edit, 1)
        row.addWidget(btn)
        row.addWidget(self.btn_import)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(6)
        lay.addLayout(row)
        lay.addWidget(self.list)

    def _insert(self, code, name, checked):
        it = QListWidgetItem(f"{code}  {name}")
        it.setData(ROLE_ID, code)
        it.setFlags(it.flags() | Qt.ItemIsUserCheckable)
        it.setCheckState(Qt.Checked if checked else Qt.Unchecked)
        self.list.addItem(it)
        return it

    def _find(self, code):
        for i in range(self.list.count()):
            if self.list.item(i).data(ROLE_ID) == code:
                return self.list.item(i)
        return None

    def _filter(self, text):
        t = " ".join(text.split()).lower()
        for i in range(self.list.count()):
            item = self.list.item(i)
            item.setHidden(bool(t) and t not in " ".join(item.text().split()).lower())

    def add(self):
        text = " ".join(self.edit.text().split())
        if not text:
            return
        code, name = split_label(text)
        if code is None:
            visible = [self.list.item(i) for i in range(self.list.count()) if not self.list.item(i).isHidden()]
            if len(visible) == 1:
                item = visible[0]
            else:
                QMessageBox.information(self, APP_NAME, "Para cadastrar uma empresa nova, digite o código e o nome, "
                                        "por exemplo: 10 EMPRESA EXEMPLO LTDA.\n\n"
                                        "Para marcar uma existente, filtre até sobrar só ela e aperte Enter.")
                return
        else:
            item = self._find(code)
            if item is None:
                if not name:
                    QMessageBox.information(self, APP_NAME, f"O código {code} não está cadastrado.\n\n"
                                            "Digite o código e o nome para cadastrar, por exemplo: "
                                            f"{int(code)} EMPRESA EXEMPLO LTDA.")
                    return
                self.pending[code] = name
                item = self._insert(code, name, True)
                self.list.sortItems()
        item.setCheckState(Qt.Checked)
        self.edit.clear()
        self.list.scrollToItem(item)

    def mark_codes(self, codes, replace=False):
        """Marca os códigos informados. Devolve (marcados, não cadastrados)."""
        if replace:
            for i in range(self.list.count()):
                self.list.item(i).setCheckState(Qt.Unchecked)
        found, missing = [], []
        for code in codes:
            item = self._find(code)
            if item:
                item.setCheckState(Qt.Checked)
                found.append(code)
            else:
                missing.append(code)
        return found, missing

    def import_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Importar códigos de empresas", "",
                                              "Planilhas (*.xlsx *.xlsm *.csv);;Todos os arquivos (*.*)")
        if not path:
            return
        try:
            dlg = ImportCodesDialog(self, path)
        except Exception as ex:
            QMessageBox.warning(self, APP_NAME, f"Não foi possível abrir a planilha:\n{ex}")
            return
        if not dlg.exec():
            return
        codes, invalid = dlg.result_codes()
        found, missing = self.mark_codes(codes, dlg.replace.isChecked())
        box = QMessageBox(self)
        box.setWindowTitle("Importação da planilha")
        box.setIcon(QMessageBox.Warning if missing or invalid else QMessageBox.Information)
        msg = f"{len(found)} empresa(s) marcada(s) a partir da planilha."
        if missing:
            msg += f"\n\n{len(missing)} código(s) não estão cadastrados: " + ", ".join(missing[:30])
            if len(missing) > 30:
                msg += f" e mais {len(missing) - 30}"
        if invalid:
            msg += f"\n\n{len(invalid)} valor(es) da coluna não são códigos válidos e foram ignorados."
        msg += "\n\nRevise a lista e clique em Salvar para gravar."
        box.setText(msg)
        details = []
        if missing:
            details.append("Códigos não cadastrados:\n" + "\n".join(missing))
        if invalid:
            details.append("Valores ignorados:\n" + "\n".join(invalid))
        if details:
            box.setDetailedText("\n\n".join(details))
        box.exec()

    def selected(self):
        return sorted(self.list.item(i).data(ROLE_ID) for i in range(self.list.count())
                      if self.list.item(i).checkState() == Qt.Checked)

    def new_companies(self):
        sel = set(self.selected())
        return [(c, n) for c, n in self.pending.items() if c in sel]


class ItemDialog(QDialog):
    def __init__(self, parent, db, item=None, folder_id=None, companies=None):
        super().__init__(parent)
        self.setWindowTitle("Editar automação" if item else "Nova automação")
        self.setMinimumWidth(560)
        item = item or {}

        self.name = QLineEdit(item.get("name", ""))
        self.name.setPlaceholderText("Ex.: Conciliação bancária diária")
        self.type = QComboBox()
        self.type.setEditable(True)
        self.type.addItems(db.all_types())
        self.type.setCurrentText(item.get("type", DEFAULT_TYPES[0]))
        self.folder = QComboBox()
        self.folder.addItem("Início", "")
        for fid, path in db.folder_paths():
            self.folder.addItem(path, fid)
        target = item.get("folder_id", folder_id) or ""
        self.folder.setCurrentIndex(max(self.folder.findData(target), 0))
        self.location = QLineEdit(item.get("location", ""))
        self.location.setPlaceholderText(r"Ex.: \\servidor\automacoes\financeiro\conciliacao.py")
        self.companies = CompanyPicker(db, item.get("companies") or companies or [])
        self.desc = QPlainTextEdit(item.get("description", ""))
        self.desc.setPlaceholderText("O que ela faz, quando roda, de onde lê e onde grava, quem é o responsável...")
        self.desc.setMinimumHeight(140)

        form = QFormLayout()
        form.setSpacing(10)
        form.addRow("Nome *", self.name)
        form.addRow("Tipo", self.type)
        form.addRow("Empresas", self.companies)
        form.addRow("Pasta", self.folder)
        form.addRow("Local", self.location)
        form.addRow("Descrição", self.desc)

        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("Salvar")
        buttons.button(QDialogButtonBox.Save).setProperty("primary", True)
        buttons.button(QDialogButtonBox.Cancel).setText("Cancelar")
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(20, 20, 20, 16)
        lay.addLayout(form)
        lay.addSpacing(8)
        lay.addWidget(buttons)

    def _accept(self):
        if self.companies.edit.hasFocus() and self.companies.edit.text().strip():
            self.companies.add()
            return
        if not self.name.text().strip():
            QMessageBox.warning(self, APP_NAME, "Informe o nome da automação.")
            self.name.setFocus()
            return
        self.accept()

    def values(self):
        return {
            "name": self.name.text().strip(),
            "type": self.type.currentText().strip(),
            "companies": self.companies.selected(),
            "new_companies": self.companies.new_companies(),
            "folder_id": self.folder.currentData() or None,
            "location": self.location.text().strip(),
            "description": self.desc.toPlainText().strip(),
        }


class HistoryDialog(QDialog):
    ACTIONS = [("Todas as ações", None), ("Criações", ("criou",)), ("Edições", ("editou",)),
               ("Renomeações", ("renomeou",)), ("Exclusões", ("excluiu",)), ("Anexos", ("anexou", "desanexou"))]
    LABELS = {"anexou": "Anexou arquivo", "desanexou": "Removeu anexo", "converteu": "Converteu empresas"}

    def __init__(self, parent, db, theme, target_id=None, target_name=None):
        super().__init__(parent)
        self.db, self.t, self.target_id = db, theme, target_id
        self.setWindowTitle(f"Histórico: {target_name}" if target_name else f"Histórico da base {db.name}")
        self.resize(1000, 620)

        self.search = QLineEdit()
        self.search.setPlaceholderText("Filtrar por nome, usuário, pasta ou conteúdo alterado...")
        self.search.setClearButtonEnabled(True)
        self.search.textChanged.connect(self._fill)
        self.action = QComboBox()
        for label, val in self.ACTIONS:
            self.action.addItem(label, val)
        self.action.currentIndexChanged.connect(self._fill)
        self.count = QLabel()

        self.table = QTreeWidget()
        self.table.setHeaderLabels(["Data e hora", "Usuário", "Ação", "Nome", "Pasta"])
        self.table.setRootIsDecorated(False)
        self.table.setUniformRowHeights(True)
        for col, w in enumerate([150, 110, 170, 240]):
            self.table.setColumnWidth(col, w)
        self.table.itemSelectionChanged.connect(self._show)

        self.detail = QTextBrowser()
        self.detail.setWordWrapMode(QTextOption.WrapAtWordBoundaryOrAnywhere)
        panel = QFrame()
        panel.setObjectName("panel")
        pl = QVBoxLayout(panel)
        pl.setContentsMargins(14, 12, 14, 12)
        pl.addWidget(self.detail)

        split = QSplitter(Qt.Vertical)
        split.addWidget(self.table)
        split.addWidget(panel)
        split.setSizes([320, 260])
        split.setChildrenCollapsible(False)

        top = QHBoxLayout()
        top.addWidget(self.search, 1)
        top.addWidget(self.action)
        top.addWidget(self.count)
        close = QPushButton("Fechar")
        close.clicked.connect(self.accept)
        bottom = QHBoxLayout()
        bottom.addStretch()
        bottom.addWidget(close)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 16, 16, 12)
        lay.addLayout(top)
        lay.addWidget(split, 1)
        lay.addLayout(bottom)
        self._fill()

    @staticmethod
    def label(h):
        return HistoryDialog.LABELS.get(h["action"]) or f'{h["action"].capitalize()} {h["kind"]}'

    def _fill(self):
        q = self.search.text().strip().lower()
        act = self.action.currentData()
        self.table.clear()
        n = 0
        for h in reversed(self.db.data.get("history", [])):
            if self.target_id and h.get("target_id") != self.target_id:
                continue
            if act and h.get("action") not in act:
                continue
            if q:
                blob = " ".join([h.get("name", ""), h.get("user", ""), h.get("path", ""), h.get("note", ""),
                                 self.label(h)] + [str(v) for c in h.get("changes", []) for v in c]).lower()
                if q not in blob:
                    continue
            row = QTreeWidgetItem([h.get("at", ""), h.get("user", ""), self.label(h),
                                   h.get("name", ""), h.get("path", "")])
            row.setData(0, ROLE_ID, h)
            self.table.addTopLevelItem(row)
            n += 1
        self.count.setText(f"{n} registro(s)")
        if n:
            self.table.setCurrentItem(self.table.topLevelItem(0))
        else:
            self.detail.setHtml(f"<p style='color:{self.t['faint']}'>Nenhum registro encontrado.</p>")

    def _show(self):
        row = self.table.currentItem()
        if not row:
            return
        h, t, e = row.data(0, ROLE_ID), self.t, html.escape
        out = [f"<div style='font-size:12pt;font-weight:700'>{e(self.label(h))}: {e(h.get('name', ''))}</div>",
               f"<p style='color:{t['muted']}'>{e(h.get('user', ''))} em {e(h.get('at', ''))}"
               f"{' &middot; ' + e(h['path']) if h.get('path') else ''}</p>"]
        if h.get("note"):
            out.append(f"<p>{e(h['note'])}</p>")
        changes = h.get("changes") or []
        if changes:
            cell = f"padding:6px 10px;border-bottom:1px solid {t['border']}"
            head = f"{cell};color:{t['muted']};font-weight:600"

            def fmt(v):
                return e(v).replace("\n", "<br>") if v else f"<span style='color:{t['faint']}'>(vazio)</span>"
            rows = "".join(
                f"<tr><td style='{cell};font-weight:600' width='110'>{e(c[0])}</td>"
                f"<td style='{cell};color:{t['danger']}'>{fmt(c[1])}</td>"
                f"<td style='{cell};color:{t['add']}'>{fmt(c[2])}</td></tr>" for c in changes)
            out.append(f"<table width='100%' cellspacing='0'><tr><td style='{head}'>Campo</td>"
                       f"<td style='{head}'>Antes</td><td style='{head}'>Depois</td></tr>{rows}</table>")
        self.detail.setHtml("".join(out))


class AssignCompaniesDialog(QDialog):
    """Marca ou desmarca as empresas de uma automação."""

    def __init__(self, parent, db, item):
        super().__init__(parent)
        self.setWindowTitle(f'Atribuir empresas: {item["name"]}')
        self.resize(620, 600)
        info = QLabel("Marque as empresas que usam esta automação, ou use \"Importar do Excel\" para marcar "
                      "pelos códigos de uma planilha. Para cadastrar uma nova, digite \"código nome\" e aperte Enter.")
        info.setWordWrap(True)
        self.picker = CompanyPicker(db, item.get("companies") or [])
        self.picker.list.setMinimumHeight(360)
        self.picker.list.setMaximumHeight(16777215)
        self.count = QLabel()
        self.count.setMinimumWidth(130)
        self.picker.list.itemChanged.connect(self._count)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("Salvar")
        buttons.button(QDialogButtonBox.Save).setProperty("primary", True)
        buttons.button(QDialogButtonBox.Cancel).setText("Cancelar")
        buttons.accepted.connect(self._accept)
        buttons.rejected.connect(self.reject)
        bottom = QHBoxLayout()
        bottom.addWidget(self.count)
        bottom.addStretch()
        bottom.addWidget(buttons)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 16, 16, 12)
        lay.addWidget(info)
        lay.addWidget(self.picker, 1)
        lay.addLayout(bottom)
        self._count()

    def _count(self, *_):
        self.count.setText(f"{len(self.picker.selected())} selecionada(s)")

    def _accept(self):
        if self.picker.edit.hasFocus() and self.picker.edit.text().strip():
            self.picker.add()
            self._count()
            return
        self.accept()

    def selected(self):
        return self.picker.selected()

    def new_companies(self):
        return self.picker.new_companies()


class AssignItemsDialog(QDialog):
    """Marca ou desmarca as automações ligadas a uma empresa, organizadas pelas pastas."""

    def __init__(self, parent, db, company):
        super().__init__(parent)
        self.db = db
        label = db.company_label(company)
        self.setWindowTitle(f"Atribuir automações: {label}")
        self.resize(640, 620)
        info = QLabel(f'Marque as automações usadas por "{label}". Marcar uma pasta marca tudo dentro dela.')
        info.setWordWrap(True)
        self.filter = QLineEdit()
        self.filter.setPlaceholderText("Filtrar automações...")
        self.filter.setClearButtonEnabled(True)
        self.filter.textChanged.connect(self._filter)
        self.tree = QTreeWidget()
        self.tree.setHeaderLabels(["Automação", "Tipo"])
        self.tree.setColumnWidth(0, 420)
        linked = {i["id"] for i in db.company_items(company)}
        self.item_nodes = []
        self.tree.blockSignals(True)
        self._add(self.tree.invisibleRootItem(), None, linked)
        self.tree.blockSignals(False)
        self.tree.expandAll()
        self.tree.itemChanged.connect(self._count)
        self.count = QLabel()
        self.count.setMinimumWidth(110)

        b_all = QPushButton("Marcar visíveis")
        b_all.clicked.connect(lambda: self._mark(Qt.Checked))
        b_none = QPushButton("Desmarcar visíveis")
        b_none.clicked.connect(lambda: self._mark(Qt.Unchecked))
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("Salvar")
        buttons.button(QDialogButtonBox.Save).setProperty("primary", True)
        buttons.button(QDialogButtonBox.Cancel).setText("Cancelar")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        bottom = QHBoxLayout()
        bottom.addWidget(b_all)
        bottom.addWidget(b_none)
        bottom.addWidget(self.count)
        bottom.addStretch()
        bottom.addWidget(buttons)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 16, 16, 12)
        lay.addWidget(info)
        lay.addWidget(self.filter)
        lay.addWidget(self.tree, 1)
        lay.addLayout(bottom)
        self._count()

    def _add(self, parent_node, fid, linked):
        style = self.style()
        for f in self.db.subfolders(fid):
            ids = Database.descendants(self.db.data, f["id"])
            if not any(i.get("folder_id") in ids for i in self.db.data["items"]):
                continue
            node = QTreeWidgetItem([f["name"], ""])
            node.setIcon(0, style.standardIcon(QStyle.SP_DirIcon))
            node.setFlags(node.flags() | Qt.ItemIsUserCheckable | Qt.ItemIsAutoTristate)
            parent_node.addChild(node)
            self._add(node, f["id"], linked)
        for it in self.db.items_in(fid):
            node = QTreeWidgetItem([it["name"], it.get("type", "")])
            node.setIcon(0, style.standardIcon(QStyle.SP_FileIcon))
            node.setFlags((node.flags() | Qt.ItemIsUserCheckable) & ~Qt.ItemIsAutoTristate)
            node.setData(0, ROLE_ID, it["id"])
            parent_node.addChild(node)
            node.setCheckState(0, Qt.Checked if it["id"] in linked else Qt.Unchecked)
            self.item_nodes.append(node)

    def _filter(self, text):
        t = text.strip().lower()

        def walk(node):
            visible = False
            for i in range(node.childCount()):
                ch = node.child(i)
                if ch.data(0, ROLE_ID):
                    show = not t or t in ch.text(0).lower() or t in ch.text(1).lower()
                else:
                    show = walk(ch)
                ch.setHidden(not show)
                visible = visible or show
            return visible
        walk(self.tree.invisibleRootItem())
        self.tree.expandAll()

    def _mark(self, state):
        for n in self.item_nodes:
            if not n.isHidden() and not (n.parent() and self._hidden_up(n.parent())):
                n.setCheckState(0, state)

    @staticmethod
    def _hidden_up(node):
        while node:
            if node.isHidden():
                return True
            node = node.parent()
        return False

    def _count(self, *_):
        self.count.setText(f"{len(self.selected())} marcada(s)")

    def selected(self):
        return [n.data(0, ROLE_ID) for n in self.item_nodes if n.checkState(0) == Qt.Checked]


class AttachList(QListWidget):
    filesDropped = Signal(list)

    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
        self.setDragDropMode(QAbstractItemView.DropOnly)
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)

    def dragEnterEvent(self, e):
        e.acceptProposedAction() if e.mimeData().hasUrls() else e.ignore()

    def dragMoveEvent(self, e):
        e.acceptProposedAction() if e.mimeData().hasUrls() else e.ignore()

    def dropEvent(self, e):
        files = [u.toLocalFile() for u in e.mimeData().urls() if u.isLocalFile() and os.path.isfile(u.toLocalFile())]
        if files:
            self.filesDropped.emit(files)
        e.acceptProposedAction()


# --------------------------------------------------------------------------- janela
class MainWindow(QMainWindow):
    def __init__(self, db, prefs):
        super().__init__()
        self.db, self.prefs = db, prefs
        self.theme_name = prefs.get("theme") if prefs.get("theme") in THEMES else "Claro"
        self.t = apply_theme(QApplication.instance(), self.theme_name)
        self.current = None          # pasta aberta (None = Início)
        self.mode = "folders"        # "folders" ou "companies"
        self.company = None          # empresa aberta (None = lista de empresas)
        self.selected = None         # ("folder"|"item"|"company", id) selecionado na lista
        self.resize(1280, 760)

        st = self.style()
        self.icon_folder = st.standardIcon(QStyle.SP_DirIcon)
        self.icon_item = st.standardIcon(QStyle.SP_FileIcon)

        self._build_menu()

        # barra superior
        self.btn_up = QPushButton("↑  Subir")
        self.btn_up.setToolTip("Subir um nível")
        self.btn_up.clicked.connect(self.go_up)
        self.crumb = QLabel()
        self.crumb.setObjectName("crumb")
        self.search = QLineEdit()
        self.search.setPlaceholderText("Buscar por nome, empresa, local, anexo...")
        self.search.setClearButtonEnabled(True)
        self.search.setFixedWidth(320)
        self.search.textChanged.connect(lambda: (self._fill_list(), self._show_details()))
        self.tabs = QTabBar()
        self.tabs.setObjectName("maintabs")
        self.tabs.setDrawBase(False)
        self.tabs.setExpanding(False)
        self.tabs.addTab("Pastas")
        self.tabs.addTab("Empresas")
        self.tabs.currentChanged.connect(self._tab_changed)
        tabrow = QFrame()
        tabrow.setObjectName("tabrow")
        trl = QHBoxLayout(tabrow)
        trl.setContentsMargins(12, 0, 12, 0)
        trl.addWidget(self.tabs)
        trl.addStretch()
        self.btn_new = QPushButton("Nova pasta")
        self.btn_new.clicked.connect(lambda: self.new_company() if self.mode == "companies" else self.new_folder())
        btn_item = QPushButton("Nova automação")
        btn_item.setProperty("primary", True)
        btn_item.clicked.connect(self.new_item)

        top = QFrame()
        top.setObjectName("topbar")
        tl = QHBoxLayout(top)
        tl.setContentsMargins(16, 10, 16, 10)
        tl.addWidget(self.btn_up)
        tl.addWidget(self.crumb, 1)
        tl.addWidget(self.search)
        tl.addWidget(self.btn_new)
        tl.addWidget(btn_item)

        # árvore de pastas
        self.tree = QTreeWidget()
        self.tree.setHeaderHidden(True)
        self.tree.itemSelectionChanged.connect(self._tree_selected)
        self.tree.setContextMenuPolicy(Qt.CustomContextMenu)
        self.tree.customContextMenuRequested.connect(self._tree_menu)

        # conteúdo da pasta
        self.listw = QTreeWidget()
        self.listw.setHeaderLabels(["Nome", "Tipo", "Empresas", "Local", "Pasta", "Atualizado em"])
        self.listw.setRootIsDecorated(False)
        self.listw.setUniformRowHeights(True)
        for col, w in enumerate([220, 130, 160, 220, 200, 120]):
            self.listw.setColumnWidth(col, w)
        self.listw.header().setStretchLastSection(False)
        self.listw.header().setSectionResizeMode(3, QHeaderView.Stretch)
        self.listw.itemSelectionChanged.connect(self._list_selected)
        self.listw.itemDoubleClicked.connect(self._list_double)
        self.listw.setContextMenuPolicy(Qt.CustomContextMenu)
        self.listw.customContextMenuRequested.connect(self._list_menu)

        # painel de detalhes
        self.details = QTextBrowser()
        self.details.setOpenLinks(False)
        self.details.anchorClicked.connect(
            lambda url: self.open_companies(url.toString()[8:]) if url.toString().startswith("empresa:") else None)
        self.details.setWordWrapMode(QTextOption.WrapAtWordBoundaryOrAnywhere)
        self.btn_assign = QPushButton("Atribuir empresas")
        self.btn_assign.clicked.connect(self.assign_selected)
        self.btn_copy = QPushButton("Copiar local")
        self.btn_copy.clicked.connect(self.copy_location)
        self.btn_hist = QPushButton("Histórico")
        self.btn_hist.clicked.connect(self.history_selected)
        self.btn_edit = QPushButton("Editar")
        self.btn_edit.clicked.connect(self.edit_selected)
        self.btn_del = QPushButton("Excluir")
        self.btn_del.setProperty("danger", True)
        self.btn_del.clicked.connect(self.delete_selected)
        self.att_box = QWidget()
        ab = QVBoxLayout(self.att_box)
        ab.setContentsMargins(0, 0, 0, 0)
        ab.setSpacing(6)
        self.att_title = QLabel("ANEXOS")
        self.att_title.setObjectName("section")
        self.att_list = AttachList()
        self.att_list.setFixedHeight(120)
        self.att_list.setToolTip("Arraste arquivos para cá para anexar")
        self.att_list.filesDropped.connect(self.attach_files)
        self.att_list.itemDoubleClicked.connect(lambda _: self.open_attachment())
        self.att_list.setContextMenuPolicy(Qt.CustomContextMenu)
        self.att_list.customContextMenuRequested.connect(self._att_menu)
        self.btn_att_add = QPushButton("Anexar")
        self.btn_att_add.clicked.connect(lambda: self.attach_files())
        self.btn_att_open = QPushButton("Abrir")
        self.btn_att_open.clicked.connect(self.open_attachment)
        self.btn_att_save = QPushButton("Salvar cópia")
        self.btn_att_save.clicked.connect(self.save_attachment)
        self.btn_att_del = QPushButton("Remover")
        self.btn_att_del.setProperty("danger", True)
        self.btn_att_del.clicked.connect(self.remove_attachment)
        self.att_list.itemSelectionChanged.connect(self._att_buttons)
        arow = QHBoxLayout()
        for b in (self.btn_att_add, self.btn_att_open, self.btn_att_save, self.btn_att_del):
            arow.addWidget(b)
        arow.addStretch()
        ab.addWidget(self.att_title)
        ab.addWidget(self.att_list)
        ab.addLayout(arow)

        panel = QFrame()
        panel.setObjectName("details")
        pl = QVBoxLayout(panel)
        pl.setContentsMargins(16, 16, 16, 16)
        pl.addWidget(self.details, 1)
        pl.addWidget(self.att_box)
        bl1 = QHBoxLayout()
        bl1.addWidget(self.btn_assign)
        bl1.addWidget(self.btn_copy)
        bl1.addStretch()
        bl = QHBoxLayout()
        bl.addStretch()
        bl.addWidget(self.btn_hist)
        bl.addWidget(self.btn_edit)
        bl.addWidget(self.btn_del)
        pl.addLayout(bl1)
        pl.addLayout(bl)

        split = QSplitter()
        split.addWidget(self.tree)
        split.addWidget(self.listw)
        split.addWidget(panel)
        split.setSizes([240, 620, 420])
        split.setChildrenCollapsible(False)

        central = QWidget()
        central.setObjectName("central")
        cl = QVBoxLayout(central)
        cl.setContentsMargins(0, 0, 0, 0)
        cl.setSpacing(0)
        cl.addWidget(tabrow)
        cl.addWidget(top)
        body = QWidget()
        bl2 = QVBoxLayout(body)
        bl2.setContentsMargins(12, 12, 12, 12)
        bl2.addWidget(split)
        cl.addWidget(body, 1)
        self.setCentralWidget(central)

        # barra de status
        self.db_label = QLabel()
        btn_db = QPushButton("Trocar base")
        btn_db.setObjectName("link")
        btn_db.clicked.connect(self.open_base_dialog)
        self.statusBar().addPermanentWidget(self.db_label)
        self.statusBar().addPermanentWidget(btn_db)

        # atalhos
        QShortcut(QKeySequence("Ctrl+N"), self, self.new_item)
        QShortcut(QKeySequence("Ctrl+Shift+N"), self, lambda: self.new_folder())
        QShortcut(QKeySequence("F2"), self, self.edit_selected)
        QShortcut(QKeySequence.Delete, self, self.delete_selected)
        QShortcut(QKeySequence("Ctrl+F"), self, self.search.setFocus)
        QShortcut(QKeySequence("Backspace"), self.listw, self.go_up)

        # sincronização com outros usuários
        self.timer = QTimer(self)
        self.timer.timeout.connect(self.poll)
        self.timer.start(2000)

        remember_base(self.prefs, self.db.path)
        self.refresh()

    # ------------------------------------------------------------- menu
    def _build_menu(self):
        mb = self.menuBar()
        m_base = mb.addMenu("Base")
        m_base.addAction("Nova base...", self.new_base_dialog)
        m_base.addAction("Abrir base...", self.open_base_dialog, QKeySequence("Ctrl+O"))
        self.m_recent = m_base.addMenu("Bases recentes")
        self.m_recent.aboutToShow.connect(self._fill_recent)
        m_base.addSeparator()
        m_base.addAction("Mostrar arquivo da base", self.show_base_path)
        m_base.addSeparator()
        m_base.addAction("Sair", self.close)

        m_view = mb.addMenu("Exibir")
        m_theme = m_view.addMenu("Tema")
        group = QActionGroup(self)
        group.setExclusive(True)
        for name in THEMES:
            act = QAction(name, self, checkable=True)
            act.setChecked(name == self.theme_name)
            act.triggered.connect(lambda _=False, n=name: self.set_theme(n))
            group.addAction(act)
            m_theme.addAction(act)

        m_comp = mb.addMenu("Empresas")
        m_comp.addAction("Ver empresas", lambda: self.open_companies(), QKeySequence("Ctrl+E"))
        m_comp.addAction("Nova empresa...", self.new_company)

        m_hist = mb.addMenu("Histórico")
        m_hist.addAction("Histórico completo da base", self.show_history, QKeySequence("Ctrl+H"))
        m_hist.addAction("Histórico do item selecionado", self.history_selected)

    def _fill_recent(self):
        self.m_recent.clear()
        recent = self.prefs.get("recent", [])
        cur = os.path.normcase(self.db.path)
        for p in recent:
            act = self.m_recent.addAction(p, lambda path=p: self.switch_base(path))
            act.setCheckable(True)
            act.setChecked(os.path.normcase(p) == cur)
        if not recent:
            self.m_recent.addAction("Nenhuma base recente").setEnabled(False)
        else:
            self.m_recent.addSeparator()
            self.m_recent.addAction("Limpar lista", self._clear_recent)

    def _clear_recent(self):
        self.prefs["recent"] = [self.db.path]
        save_prefs(self.prefs)

    # ------------------------------------------------------------- base e tema
    def switch_base(self, path):
        if os.path.normcase(os.path.abspath(path)) == os.path.normcase(self.db.path):
            return
        try:
            db = Database(path)
        except Exception as ex:
            QMessageBox.warning(self, APP_NAME, f"Não foi possível abrir a base:\n{path}\n\n{ex}")
            return
        self.db = db
        remember_base(self.prefs, db.path)
        self.current, self.selected, self.mode, self.company = None, None, "folders", None
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self.refresh()
        self.statusBar().showMessage(f'Base "{db.name}" aberta.', 4000)

    def open_base_dialog(self):
        path = ask_open_base(self)
        if path:
            self.switch_base(path)

    def new_base_dialog(self):
        path = ask_new_base(self)
        if path:
            self.switch_base(path)

    def show_base_path(self):
        QGuiApplication.clipboard().setText(self.db.path)
        QMessageBox.information(self, APP_NAME, f"A base atual fica em:\n\n{self.db.path}\n\n"
                                "O caminho foi copiado para a área de transferência.")

    def set_theme(self, name):
        self.theme_name = name
        self.t = apply_theme(QApplication.instance(), name)
        self.prefs["theme"] = name
        save_prefs(self.prefs)
        self._show_details()

    # ------------------------------------------------------------- montagem
    def _real_company(self, name):
        """Código da empresa se ela ainda existe, senão None."""
        if name == NO_COMPANY:
            return NO_COMPANY
        return name if name and self.db.company(name) else None

    def refresh(self):
        if self.current and not self.db.folder(self.current):
            self.current = None
        if self.company:
            self.company = self._real_company(self.company)
        if self.selected and self.selected[0] == "company":
            real = self._real_company(self.selected[1])
            self.selected = ("company", real) if real else None
        self.btn_new.setText("Nova empresa" if self.mode == "companies" else "Nova pasta")
        self.tabs.blockSignals(True)
        self.tabs.setCurrentIndex(1 if self.mode == "companies" else 0)
        self.tabs.blockSignals(False)
        self.setWindowTitle(f"{APP_NAME}  |  {self.db.name}")
        self.db_label.setText(f"Base: {self.db.path}   |   Usuário: {current_user()}")
        self._fill_tree()
        self._fill_list()
        self._show_details()

    def _fill_tree(self):
        expanded = getattr(self, "_expanded", {("folder", ""): True, ("companies", ""): True})

        def collect(node):
            for i in range(node.childCount()):
                ch = node.child(i)
                expanded[(ch.data(0, ROLE_KIND), ch.data(0, ROLE_ID))] = ch.isExpanded()
                collect(ch)
        collect(self.tree.invisibleRootItem())
        self._expanded = expanded

        def node(text, kind, oid, icon):
            n = QTreeWidgetItem([text])
            n.setIcon(0, icon)
            n.setData(0, ROLE_ID, oid)
            n.setData(0, ROLE_KIND, kind)
            return n

        self.tree.blockSignals(True)
        self.tree.clear()
        st = self.style()
        if self.mode == "companies":
            self._fill_company_tree(node, expanded)
            self.tree.blockSignals(False)
            return
        root = node("Início", "folder", "", st.standardIcon(QStyle.SP_DirHomeIcon))
        self.tree.addTopLevelItem(root)
        to_select = root

        def add(parent_node, parent_id):
            nonlocal to_select
            for f in self.db.subfolders(parent_id):
                n = node(f["name"], "folder", f["id"], self.icon_folder)
                parent_node.addChild(n)
                if expanded.get(("folder", f["id"])):
                    n.setExpanded(True)
                if self.mode == "folders" and f["id"] == self.current:
                    to_select = n
                add(n, f["id"])
        add(root, None)
        root.setExpanded(expanded.get(("folder", ""), True))
        p = to_select.parent()
        while p:
            p.setExpanded(True)
            p = p.parent()
        self.tree.setCurrentItem(to_select)
        self.tree.blockSignals(False)

    def _fill_company_tree(self, node, expanded):
        croot = node("Empresas", "companies", "", self.style().standardIcon(QStyle.SP_DirHomeIcon))
        self.tree.addTopLevelItem(croot)
        to_select = croot
        counts = {}
        for i in self.db.data["items"]:
            for code in i.get("companies") or []:
                counts[code] = counts.get(code, 0) + 1
        for c in self.db.companies():
            n = node(f'{c["code"]}  {c["name"]}  ({counts.get(c["code"], 0)})', "company", c["code"], self.icon_folder)
            croot.addChild(n)
            if self.company == c["code"]:
                to_select = n
        n_none = len(self.db.company_items(None))
        if n_none:
            n = node(f"Sem empresa  ({n_none})", "company", NO_COMPANY, self.icon_folder)
            f = n.font(0)
            f.setItalic(True)
            n.setFont(0, f)
            croot.addChild(n)
            if self.company == NO_COMPANY:
                to_select = n
        croot.setExpanded(True)
        self.tree.setCurrentItem(to_select)

    def _fill_list(self):
        q = self.search.text().strip().lower()
        searching = bool(q)
        self.listw.blockSignals(True)
        self.listw.clear()
        self.listw.setColumnHidden(4, not searching)

        companies = []
        if searching:
            folders = [f for f in self.db.data["folders"] if q in f["name"].lower()]
            items = [i for i in self.db.data["items"] if q in " ".join([
                i.get("name", ""), i.get("type", ""), i.get("location", ""),
                i.get("description", ""), self.db._labels(i.get("companies")),
                " ".join(a["file"] for a in i.get("attachments") or []), self.db.folder_path(i.get("folder_id"))]).lower()]
            folders.sort(key=lambda f: f["name"].lower())
            items.sort(key=lambda i: i["name"].lower())
        elif self.mode == "companies":
            folders = []
            if self.company:
                items = self.db.company_items(None if self.company == NO_COMPANY else self.company)
                self.listw.setColumnHidden(4, False)
            else:
                items = []
                counts = {}
                for i in self.db.data["items"]:
                    for code in i.get("companies") or []:
                        counts[code] = counts.get(code, 0) + 1
                companies = [(c["code"], counts.get(c["code"], 0)) for c in self.db.companies()]
                n_none = len(self.db.company_items(None))
                if n_none:
                    companies.append((NO_COMPANY, n_none))
        else:
            folders = self.db.subfolders(self.current)
            items = self.db.items_in(self.current)

        to_select = None
        for c, n in companies:
            row = QTreeWidgetItem(["Sem empresa" if c == NO_COMPANY else self.db.company_label(c), "Empresa",
                                   f"{n} automação(ões)", "", "", ""])
            row.setIcon(0, self.icon_folder)
            row.setData(0, ROLE_ID, c)
            row.setData(0, ROLE_KIND, "company")
            if c == NO_COMPANY:
                f = row.font(0)
                f.setItalic(True)
                row.setFont(0, f)
            self.listw.addTopLevelItem(row)
            if self.selected == ("company", c):
                to_select = row
        for f in folders:
            row = QTreeWidgetItem([f["name"], "Pasta", "", "", self.db.folder_path(f.get("parent_id")), ""])
            row.setIcon(0, self.icon_folder)
            row.setData(0, ROLE_ID, f["id"])
            row.setData(0, ROLE_KIND, "folder")
            self.listw.addTopLevelItem(row)
            if self.selected == ("folder", f["id"]):
                to_select = row
        for i in items:
            codes = i.get("companies") or []
            comp = self.db.company_label(codes[0]) if len(codes) == 1 else ", ".join(codes)
            row = QTreeWidgetItem([i["name"], i.get("type", ""), comp, i.get("location", ""),
                                   self.db.folder_path(i.get("folder_id")), i.get("updated_at", "")])
            row.setIcon(0, self.icon_item)
            row.setToolTip(2, "\n".join(self.db.company_label(c) for c in codes))
            row.setToolTip(3, i.get("location", ""))
            n_att = len(i.get("attachments") or [])
            if n_att:
                row.setToolTip(0, f"{n_att} anexo(s)")
            row.setData(0, ROLE_ID, i["id"])
            row.setData(0, ROLE_KIND, "item")
            self.listw.addTopLevelItem(row)
            if self.selected == ("item", i["id"]):
                to_select = row

        if not folders and not items and not companies:
            if searching:
                msg = "Nenhum resultado para a busca."
            elif self.mode == "companies" and self.company == NO_COMPANY:
                msg = "Todas as automações têm pelo menos uma empresa vinculada."
            elif self.mode == "companies" and self.company:
                msg = "Nenhuma automação ligada a esta empresa. Use \"Atribuir automações\" no painel ao lado."
            elif self.mode == "companies":
                msg = "Nenhuma empresa cadastrada. Use \"Nova empresa\" para começar."
            else:
                msg = "Pasta vazia. Use \"Nova pasta\" ou \"Nova automação\" para começar."
            empty = QTreeWidgetItem([msg])
            empty.setFlags(Qt.NoItemFlags)
            self.listw.addTopLevelItem(empty)
            empty.setFirstColumnSpanned(True)

        if to_select:
            self.listw.setCurrentItem(to_select)
        else:
            self.selected = None
        self.listw.blockSignals(False)

        if searching:
            self.crumb.setText(f"Resultados da busca: {len(folders) + len(items)}")
            self.btn_up.setEnabled(False)
        elif self.mode == "companies":
            name = "Sem empresa" if self.company == NO_COMPANY else self.db.company_label(self.company)
            self.crumb.setText(f"Empresas / {name}" if self.company else "Empresas")
            self.btn_up.setEnabled(bool(self.company))
        else:
            self.crumb.setText(self.db.folder_path(self.current))
            self.btn_up.setEnabled(self.current is not None)

    def _last_change(self, oid):
        return next((h for h in reversed(self.db.data.get("history", [])) if h.get("target_id") == oid), None)

    def _show_details(self):
        e, t = html.escape, self.t
        label = f"color:{t['muted']};font-size:8pt;font-weight:600;letter-spacing:1px;margin-top:14px"
        meta_style = f"color:{t['faint']};font-size:8pt;margin-top:18px"
        empty = f"<i style='color:{t['faint']}'>"
        if self.selected:
            kind, oid = self.selected
        elif self.mode == "companies":
            kind, oid = ("company", self.company) if self.company else ("companies", None)
        else:
            kind, oid = "folder", self.current
        self.btn_assign.setVisible(kind in ("item", "company"))
        self.btn_hist.setVisible(kind != "companies")

        if kind in ("company", "companies"):
            self.btn_copy.setVisible(False)
            self.att_box.setVisible(False)
            if kind == "companies":
                comps = self.db.companies()
                n_none = len(self.db.company_items(None))
                self.details.setHtml(f"""
                    <div style='font-size:15pt;font-weight:700'>Empresas</div>
                    <p style='{label}'>RESUMO</p>
                    <p>{len(comps)} empresa(s) cadastrada(s)<br>{n_none} automação(ões) sem empresa vinculada</p>
                    <p style='{meta_style}'>Abra uma empresa para ver as automações ligadas a ela.</p>""")
                self.btn_edit.setEnabled(False)
                self.btn_del.setEnabled(False)
                return
            items = self.db.company_items(None if oid == NO_COMPANY else oid)
            rows = "".join(f"<tr><td style='padding:3px 10px 3px 0'>{e(i['name'])}</td>"
                           f"<td style='padding:3px 0;color:{t['muted']}'>{e(self.db.folder_path(i.get('folder_id')))}</td></tr>"
                           for i in items) or f"<tr><td>{empty}Nenhuma automação ligada</i></td></tr>"
            types = {}
            for i in items:
                types[i.get("type") or "Sem tipo"] = types.get(i.get("type") or "Sem tipo", 0) + 1
            tline = ", ".join(f"{e(k)} ({v})" for k, v in sorted(types.items(), key=lambda kv: -kv[1]))
            title = "Sem empresa vinculada" if oid == NO_COMPANY else self.db.company_label(oid)
            self.details.setHtml(f"""
                <div style='font-size:15pt;font-weight:700'>{e(title)}</div>
                <p style='{label}'>AUTOMAÇÕES ({len(items)})</p><table cellspacing='0'>{rows}</table>
                {f"<p style='{label}'>TIPOS</p><p>{tline}</p>" if tline else ""}""")
            real = oid != NO_COMPANY
            self.btn_assign.setVisible(real)
            self.btn_assign.setText("Atribuir automações")
            self.btn_hist.setEnabled(real)
            self.btn_edit.setEnabled(real)
            self.btn_del.setEnabled(real)
            return

        self.btn_assign.setText("Atribuir empresas")
        if kind == "item" and self.db.item(oid):
            it = self.db.item(oid)
            desc = e(it.get("description", "")).replace("\n", "<br>") or f"{empty}Sem descrição</i>"
            loc = e(it.get("location", "")).replace("\\\\", "\\\u2060\\") or f"{empty}Não informado</i>"
            meta = f"Criado por {e(it.get('created_by', '?'))} em {e(it.get('created_at', '?'))}<br>" \
                   f"Atualizado por {e(it.get('updated_by', '?'))} em {e(it.get('updated_at', '?'))}"
            n_hist = sum(1 for h in self.db.data.get("history", []) if h.get("target_id") == oid)
            badge = f"background:{t['hover']};color:{t['text']}"
            comps = "<br>".join(f"<a href='empresa:{e(c)}' style='text-decoration:none;{badge}'>&nbsp;"
                                f"{e(self.db.company_label(c))}&nbsp;</a>"
                                for c in it.get("companies") or []) or f"{empty}Nenhuma empresa</i>"
            self._fill_attachments(it)
            self.details.setHtml(f"""
                <div style='font-size:15pt;font-weight:700'>{e(it['name'])}</div>
                <div style='margin-top:6px'><span style='background:{t['sel']};color:{t['sel_text']};font-weight:600'>
                &nbsp;{e(it.get('type', '') or 'Sem tipo')}&nbsp;</span></div>
                <p style='{label}'>EMPRESAS</p><p>{comps}</p>
                <p style='{label}'>PASTA</p><p>{e(self.db.folder_path(it.get('folder_id')))}</p>
                <p style='{label}'>LOCAL</p><p style='font-family:Consolas,monospace'>{loc}</p>
                <p style='{label}'>DESCRIÇÃO</p><p>{desc}</p>
                <p style='{meta_style}'>{meta}<br>{n_hist} registro(s) no histórico</p>""")
            self.att_box.setVisible(True)
            self.btn_copy.setVisible(bool(it.get("location")))
            for b in (self.btn_hist, self.btn_edit, self.btn_del):
                b.setEnabled(True)
            return

        fid = oid if kind == "folder" else self.current
        ids = Database.descendants(self.db.data, fid) if fid else {None} | {f["id"] for f in self.db.data["folders"]}
        n_items = sum(1 for i in self.db.data["items"] if i.get("folder_id") in ids)
        n_sub = len(self.db.subfolders(fid))
        f = self.db.folder(fid)
        extra = f"<p style='{meta_style}'>Criada por {e(f.get('created_by', '?'))} em {e(f.get('created_at', '?'))}</p>" if f and f.get("created_by") else ""
        self.details.setHtml(f"""
            <div style='font-size:15pt;font-weight:700'>{e(f['name']) if f else 'Início'}</div>
            <p style='{label}'>CAMINHO</p><p>{e(self.db.folder_path(fid))}</p>
            <p style='{label}'>CONTEÚDO</p>
            <p>{n_sub} subpasta(s) direta(s)<br>{n_items} automação(ões) no total, incluindo subpastas</p>{extra}""")
        self.btn_copy.setVisible(False)
        self.att_box.setVisible(False)
        self.btn_hist.setEnabled(True)
        self.btn_edit.setEnabled(f is not None)
        self.btn_del.setEnabled(f is not None)

    # ------------------------------------------------------------- navegação
    def _tree_selected(self):
        node = self.tree.currentItem()
        if not node:
            return
        kind, oid = node.data(0, ROLE_KIND), node.data(0, ROLE_ID)
        if kind == "folder":
            self.mode, self.current = "folders", oid or None
        else:
            self.mode, self.company = "companies", (oid if kind == "company" else None)
        self.btn_new.setText("Nova empresa" if self.mode == "companies" else "Nova pasta")
        self.selected = None
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self._fill_list()
        self._show_details()

    def _list_selected(self):
        row = self.listw.currentItem()
        self.selected = (row.data(0, ROLE_KIND), row.data(0, ROLE_ID)) if row and row.data(0, ROLE_KIND) else None
        self._show_details()

    def _list_double(self, row):
        kind, oid = row.data(0, ROLE_KIND), row.data(0, ROLE_ID)
        if kind == "folder":
            self.open_folder(oid)
        elif kind == "company":
            self.open_companies(oid)
        elif kind == "item":
            self.edit_item(oid)

    def _tab_changed(self, index):
        self.mode = "companies" if index == 1 else "folders"
        self.selected = None
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self.refresh()

    def open_companies(self, name=None):
        self.mode = "companies"
        self.company = self._real_company(name) if name else None
        self.selected = None
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self.refresh()

    def open_folder(self, fid):
        self.mode = "folders"
        self.current = fid
        self.selected = None
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self.refresh()

    def go_up(self):
        if self.search.text().strip():
            return
        if self.mode == "companies":
            if self.company:
                self.selected = ("company", self.company)
                self.company = None
                self.refresh()
            return
        if self.current:
            child = self.current
            f = self.db.folder(self.current)
            self.current = f.get("parent_id") if f else None
            self.selected = ("folder", child)
            self.refresh()

    # ------------------------------------------------------------- menus de contexto
    def _tree_menu(self, pos):
        node = self.tree.itemAt(pos)
        if not node:
            return
        self.tree.setCurrentItem(node)
        kind, fid = node.data(0, ROLE_KIND), node.data(0, ROLE_ID) or None
        m = QMenu(self)
        if kind in ("companies", "company"):
            m.addAction("Nova empresa", self.new_company)
            if kind == "company" and fid != NO_COMPANY:
                m.addSeparator()
                m.addAction("Atribuir automações...", lambda: self.assign_items(fid))
                m.addAction("Nova automação para esta empresa", self.new_item)
                m.addAction("Editar empresa", lambda: self.rename_company(fid))
                m.addAction("Histórico", lambda: self.company_history(fid))
                m.addAction("Excluir empresa", lambda: self.delete_company(fid))
            m.exec(self.tree.viewport().mapToGlobal(pos))
            return
        m.addAction("Nova pasta aqui", lambda: self.new_folder(fid))
        m.addAction("Nova automação aqui", self.new_item)
        if fid:
            m.addSeparator()
            m.addAction("Renomear", lambda: self.rename_folder(fid))
            m.addAction("Histórico", lambda: self.show_history(fid))
            m.addAction("Excluir", lambda: self.delete_folder(fid))
        m.exec(self.tree.viewport().mapToGlobal(pos))

    def _list_menu(self, pos):
        row = self.listw.itemAt(pos)
        m = QMenu(self)
        if row and row.data(0, ROLE_KIND):
            kind, oid = row.data(0, ROLE_KIND), row.data(0, ROLE_ID)
            if kind == "company":
                m.addAction("Abrir", lambda: self.open_companies(oid))
                if oid != NO_COMPANY:
                    m.addAction("Atribuir automações...", lambda: self.assign_items(oid))
                    m.addAction("Editar empresa", lambda: self.rename_company(oid))
                    m.addSeparator()
                    m.addAction("Excluir empresa", lambda: self.delete_company(oid))
                m.exec(self.listw.viewport().mapToGlobal(pos))
                return
            if kind == "folder":
                m.addAction("Abrir", lambda: self.open_folder(oid))
                m.addAction("Renomear", lambda: self.rename_folder(oid))
            else:
                m.addAction("Editar", lambda: self.edit_item(oid))
                m.addAction("Atribuir empresas...", lambda: self.assign_companies(oid))
                if self.mode == "companies" and self.company and self.company != NO_COMPANY \
                        and not self.search.text().strip():
                    m.addAction(f'Remover de "{self.db.company_label(self.company)}"',
                                lambda: self.unlink(oid, self.company))
                m.addAction("Anexar arquivos...", lambda: self.attach_files())
                loc = (self.db.item(oid) or {}).get("location")
                if loc:
                    m.addAction("Copiar local", lambda: QGuiApplication.clipboard().setText(loc))
            m.addAction("Histórico", lambda: self.show_history(oid))
            m.addSeparator()
            m.addAction("Excluir", lambda: (self.delete_folder if kind == "folder" else self.delete_item)(oid))
        elif self.mode == "companies" and not self.search.text().strip():
            m.addAction("Nova empresa", self.new_company)
            if self.company and self.company != NO_COMPANY:
                m.addAction("Atribuir automações...", lambda: self.assign_items(self.company))
                m.addAction("Nova automação para esta empresa", self.new_item)
        else:
            m.addAction("Nova pasta", lambda: self.new_folder())
            m.addAction("Nova automação", self.new_item)
        m.exec(self.listw.viewport().mapToGlobal(pos))

    # ------------------------------------------------------------- ações
    def _run(self, fn):
        try:
            return fn()
        except Exception as ex:
            QMessageBox.warning(self, APP_NAME, str(ex))
            return None
        finally:
            self.refresh()

    def new_folder(self, parent="__current__"):
        parent = self.current if parent == "__current__" else parent
        name, ok = QInputDialog.getText(self, "Nova pasta", f"Nome da pasta em {self.db.folder_path(parent)}:")
        if ok and name.strip():
            fid = self._run(lambda: self.db.add_folder(name.strip(), parent))
            if fid:
                self.selected = ("folder", fid) if parent == self.current else None
                self._fill_list()
                self._show_details()

    def rename_folder(self, fid):
        f = self.db.folder(fid)
        if not f:
            return
        name, ok = QInputDialog.getText(self, "Renomear pasta", "Novo nome:", text=f["name"])
        if ok and name.strip() and name.strip() != f["name"]:
            self._run(lambda: self.db.rename_folder(fid, name.strip()))

    def delete_folder(self, fid):
        f = self.db.folder(fid)
        if not f:
            return
        ids = Database.descendants(self.db.data, fid)
        n_items = sum(1 for i in self.db.data["items"] if i.get("folder_id") in ids)
        msg = f'Excluir a pasta "{f["name"]}"?'
        if len(ids) > 1 or n_items:
            msg += f"\n\nIsso também exclui {len(ids) - 1} subpasta(s) e {n_items} automação(ões) dentro dela."
        if QMessageBox.question(self, APP_NAME, msg) != QMessageBox.Yes:
            return
        if self.current in ids:
            self.current = f.get("parent_id")
        self.selected = None
        self._run(lambda: self.db.delete_folder(fid))

    def new_item(self):
        in_company = self.mode == "companies" and self.company and self.company != NO_COMPANY
        dlg = ItemDialog(self, self.db, folder_id=self.current, companies=[self.company] if in_company else None)
        if dlg.exec():
            vals = dlg.values()
            new = vals.pop("new_companies", [])
            iid = self._run(lambda: self.db.save_item(vals, new_companies=new))
            if iid:
                if self.mode == "folders":
                    self.open_folder(vals["folder_id"])
                self.selected = ("item", iid)
                self.refresh()

    def edit_item(self, iid):
        it = self.db.item(iid)
        if not it:
            return
        dlg = ItemDialog(self, self.db, item=it)
        if dlg.exec():
            vals = dlg.values()
            new = vals.pop("new_companies", [])
            self._run(lambda: self.db.save_item(vals, iid, new_companies=new))

    def delete_item(self, iid):
        it = self.db.item(iid)
        if it and QMessageBox.question(self, APP_NAME, f'Excluir a automação "{it["name"]}"?') == QMessageBox.Yes:
            self.selected = None
            self._run(lambda: self.db.delete_item(iid))

    def _target(self):
        if self.selected:
            return self.selected
        if self.mode == "companies":
            return ("company", self.company) if self.company else None
        if self.current:
            return ("folder", self.current)
        return None

    def edit_selected(self):
        t = self._target()
        if t and t[1] != NO_COMPANY:
            {"folder": self.rename_folder, "item": self.edit_item, "company": self.rename_company}[t[0]](t[1])

    def delete_selected(self):
        if self.search.hasFocus():
            return
        t = self._target()
        if t and t[1] != NO_COMPANY:
            {"folder": self.delete_folder, "item": self.delete_item, "company": self.delete_company}[t[0]](t[1])

    def assign_selected(self):
        t = self._target()
        if t and t[0] == "item":
            self.assign_companies(t[1])
        elif t and t[0] == "company" and t[1] != NO_COMPANY:
            self.assign_items(t[1])

    def history_selected(self):
        t = self._target()
        if t and t[0] == "company":
            self.company_history(t[1])
        else:
            self.show_history(t[1] if t else None)

    # ------------------------------------------------------------- empresas
    def new_company(self):
        dlg = CompanyDialog(self)
        if not dlg.exec():
            return
        code, name = dlg.values()
        if self._run(lambda: self.db.add_company(code, name)):
            if self.mode != "companies":
                self.open_companies()
            if not self.company:
                self.selected = ("company", code)
            self.refresh()

    def rename_company(self, code):
        c = self.db.company(code)
        if not c:
            return
        dlg = CompanyDialog(self, c["code"], c["name"])
        if not dlg.exec():
            return
        new_code, new_name = dlg.values()
        if new_code != code:
            n = len(self.db.company_items(code))
            if n and QMessageBox.question(
                    self, APP_NAME, f"Trocar o código de {code} para {new_code}?\n\n"
                    f"As {n} automação(ões) ligadas passam a usar o código novo.") != QMessageBox.Yes:
                return
        res = self._run(lambda: self.db.edit_company(code, new_code, new_name))
        if res:
            if self.company == code:
                self.company = res
            if self.selected == ("company", code):
                self.selected = ("company", res)
            self.refresh()
            self.statusBar().showMessage("Empresa atualizada.", 3000)

    def delete_company(self, code):
        n = len(self.db.company_items(code))
        msg = f'Excluir a empresa "{self.db.company_label(code)}"?'
        if n:
            msg += f"\n\nEla será retirada de {n} automação(ões). As automações não são excluídas."
        if QMessageBox.question(self, APP_NAME, msg) != QMessageBox.Yes:
            return
        if self.company == code:
            self.company = None
        self.selected = None
        self._run(lambda: self.db.remove_company(code))

    def assign_companies(self, iid):
        it = self.db.item(iid)
        if not it:
            return
        dlg = AssignCompaniesDialog(self, self.db, it)
        if dlg.exec():
            codes, new = dlg.selected(), dlg.new_companies()
            if self._run(lambda: self.db.set_item_companies(iid, codes, new)) is not None:
                self.statusBar().showMessage("Empresas atualizadas.", 3000)

    def assign_items(self, code):
        if not self.db.data["items"]:
            QMessageBox.information(self, APP_NAME, "Ainda não há automações cadastradas.")
            return
        dlg = AssignItemsDialog(self, self.db, code)
        if dlg.exec():
            ids = dlg.selected()
            n = self._run(lambda: self.db.set_company_items(code, ids))
            if n is not None:
                self.statusBar().showMessage(f"{n} automação(ões) alterada(s) em {self.db.company_label(code)}.", 4000)

    def unlink(self, iid, code):
        it = self.db.item(iid)
        if it:
            rest = [c for c in it.get("companies") or [] if c != code]
            self.selected = None
            self._run(lambda: self.db.set_item_companies(iid, rest))

    def company_history(self, code):
        dlg = HistoryDialog(self, self.db, self.t, None, None)
        dlg.setWindowTitle(f"Histórico: {self.db.company_label(code)}")
        dlg.search.setText(code)
        dlg.exec()

    def show_history(self, oid=None):
        name = None
        if oid:
            obj = self.db.item(oid) or self.db.folder(oid)
            name = obj["name"] if obj else None
        HistoryDialog(self, self.db, self.t, oid, name).exec()

    def _fill_attachments(self, it):
        keep = {i.data(ROLE_ID) for i in self.att_list.selectedItems()}
        self.att_list.clear()
        icon = self.style().standardIcon(QStyle.SP_FileIcon)
        for a in sorted(it.get("attachments") or [], key=lambda a: a["file"].lower()):
            row = QListWidgetItem(icon, f'{a["file"]}   ({human_size(a.get("size", 0))})')
            row.setData(ROLE_ID, a["id"])
            row.setToolTip(f'Anexado por {a.get("added_by", "?")} em {a.get("added_at", "?")}')
            self.att_list.addItem(row)
            row.setSelected(a["id"] in keep)
        n = self.att_list.count()
        self.att_title.setText(f"ANEXOS ({n})" if n else "ANEXOS  (arraste arquivos aqui)")
        self._att_buttons()

    def _att_buttons(self):
        has = bool(self.att_list.selectedItems())
        self.btn_att_open.setEnabled(has)
        self.btn_att_save.setEnabled(len(self.att_list.selectedItems()) == 1)
        self.btn_att_del.setEnabled(has)

    def _sel_item(self):
        return self.selected[1] if self.selected and self.selected[0] == "item" else None

    def _sel_attachments(self):
        it = self.db.item(self._sel_item()) or {}
        ids = {r.data(ROLE_ID) for r in self.att_list.selectedItems()}
        return it, [a for a in it.get("attachments") or [] if a["id"] in ids]

    def _att_menu(self, pos):
        row = self.att_list.itemAt(pos)
        if row and not row.isSelected():
            self.att_list.clearSelection()
            row.setSelected(True)
        m = QMenu(self)
        if self.att_list.selectedItems():
            m.addAction("Abrir", self.open_attachment)
            if len(self.att_list.selectedItems()) == 1:
                m.addAction("Salvar cópia como...", self.save_attachment)
            m.addAction("Mostrar na pasta", self.reveal_attachment)
            m.addSeparator()
            m.addAction("Remover", self.remove_attachment)
        else:
            m.addAction("Anexar arquivos...", lambda: self.attach_files())
        m.exec(self.att_list.viewport().mapToGlobal(pos))

    def attach_files(self, paths=None):
        iid = self._sel_item()
        if not iid:
            return
        if not paths:
            paths, _ = QFileDialog.getOpenFileNames(self, "Anexar arquivos", "",
                                                    "PDF (*.pdf);;Todos os arquivos (*.*)")
        if not paths:
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            n = self.db.add_attachments(iid, paths)
            self.statusBar().showMessage(f"{n} arquivo(s) anexado(s).", 4000)
        except Exception as ex:
            QApplication.restoreOverrideCursor()
            QMessageBox.warning(self, APP_NAME, str(ex))
        finally:
            if QApplication.overrideCursor():
                QApplication.restoreOverrideCursor()
            self.refresh()

    def open_attachment(self):
        it, atts = self._sel_attachments()
        for a in atts:
            path = self.db.attachment_path(it["id"], a)
            if not os.path.exists(path):
                QMessageBox.warning(self, APP_NAME, f'O arquivo "{a["file"]}" não foi encontrado na pasta de anexos.')
                continue
            QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def reveal_attachment(self):
        it, atts = self._sel_attachments()
        if atts:
            QDesktopServices.openUrl(QUrl.fromLocalFile(os.path.dirname(self.db.attachment_path(it["id"], atts[0]))))

    def save_attachment(self):
        it, atts = self._sel_attachments()
        if len(atts) != 1:
            return
        src = self.db.attachment_path(it["id"], atts[0])
        dst, _ = QFileDialog.getSaveFileName(self, "Salvar cópia", atts[0]["file"])
        if dst:
            try:
                shutil.copy2(src, dst)
                self.statusBar().showMessage("Cópia salva.", 3000)
            except OSError as ex:
                QMessageBox.warning(self, APP_NAME, f"Não foi possível salvar a cópia:\n{ex}")

    def remove_attachment(self):
        it, atts = self._sel_attachments()
        if not atts:
            return
        names = "\n".join(a["file"] for a in atts)
        if QMessageBox.question(self, APP_NAME, f"Remover {len(atts)} anexo(s)?\n\n{names}") != QMessageBox.Yes:
            return
        for a in atts:
            try:
                self.db.remove_attachment(it["id"], a["id"])
            except Exception as ex:
                QMessageBox.warning(self, APP_NAME, str(ex))
                break
        self.refresh()

    def goto_item(self, iid):
        it = self.db.item(iid)
        if not it:
            return
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self.mode = "folders"
        self.current = it.get("folder_id")
        self.selected = ("item", iid)
        self.refresh()
        self.listw.setFocus()

    def copy_location(self):
        if self.selected and self.selected[0] == "item":
            loc = (self.db.item(self.selected[1]) or {}).get("location", "")
            QGuiApplication.clipboard().setText(loc)
            self.statusBar().showMessage("Local copiado para a área de transferência.", 3000)

    def poll(self):
        try:
            if self.db.changed_on_disk():
                self.db.load()
                self.refresh()
                self.statusBar().showMessage("Base atualizada com alterações de outros usuários.", 4000)
        except Exception as ex:
            self.statusBar().showMessage(f"Sem acesso à base: {ex}", 4000)


# --------------------------------------------------------------------------- início
def main():
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setStyle("Fusion")
    if os.path.exists(resource(ICON_FILE)):
        app.setWindowIcon(QIcon(resource(ICON_FILE)))
    if sys.platform == "win32":
        try:  # faz a barra de tarefas usar o ícone do app mesmo rodando como .py
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("CatalogoAutomacoes")
        except Exception:
            pass
    prefs = load_json(PREFS_PATH)
    apply_theme(app, prefs.get("theme", "Claro"))

    candidates = [p for p in (prefs.get("db_path"), load_json(CONFIG_PATH).get("db_path")) if p]
    db = None
    for path in candidates:
        try:
            db = Database(path) if os.path.exists(path) else None
        except Exception:
            db = None
        if db:
            break
    while db is None:
        path = ask_base_startup()
        if not path:
            return
        try:
            db = Database(path)
        except Exception as ex:
            QMessageBox.warning(None, APP_NAME, f"Não foi possível abrir a base:\n{path}\n\n{ex}")

    win = MainWindow(db, prefs)
    win.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()